"""
Script para crear plantilla Excel de carga de datos de huella de carbono.
Genera un archivo con instrucciones, validaciones y ejemplos.
"""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime

# Definir estructura de la plantilla
COLUMNS = {
    'entity_id': 'ID de la entidad/instalación (ej: FACTORY001, OFFICE_HQ)',
    'scope': 'Alcance GHG (1, 2 o 3)',
    'category': 'Categoría de emisión',
    'subcategory': 'Subcategoría (opcional)',
    'activity_value': 'Valor numérico de la actividad',
    'activity_unit': 'Unidad de medida (litres, kWh, kg, m³, km, etc.)',
    'geography': 'Código país ISO-3 (GBR, USA, ESP, etc.)',
    'year': 'Año (ej: 2025)',
    'month': 'Mes (1-12, opcional)',
    'facility': 'Nombre de la instalación (opcional)',
    'description': 'Descripción de la actividad (opcional)'
}

# Datos de ejemplo
EXAMPLE_DATA = [
    {
        'entity_id': 'FACTORY001',
        'scope': 1,
        'category': 'mobile_combustion',
        'subcategory': 'diesel',
        'activity_value': 1000,
        'activity_unit': 'litres',
        'geography': 'GBR',
        'year': 2025,
        'month': 1,
        'facility': 'Plant_North',
        'description': 'Diesel para vehículos de la flota corporativa'
    },
    {
        'entity_id': 'FACTORY001',
        'scope': 2,
        'category': 'purchased_electricity',
        'subcategory': None,
        'activity_value': 50000,
        'activity_unit': 'kWh',
        'geography': 'GBR',
        'year': 2025,
        'month': 1,
        'facility': 'Plant_North',
        'description': 'Consumo eléctrico mensual de planta de producción'
    },
    {
        'entity_id': 'OFFICE_HQ',
        'scope': 3,
        'category': 'business_travel',
        'subcategory': 'flights',
        'activity_value': 8500,
        'activity_unit': 'km',
        'geography': 'USA',
        'year': 2025,
        'month': 1,
        'facility': 'Headquarters',
        'description': 'Vuelos corporativos del personal ejecutivo'
    },
    {
        'entity_id': 'WAREHOUSE_A',
        'scope': 1,
        'category': 'stationary_combustion',
        'subcategory': 'natural_gas',
        'activity_value': 5000,
        'activity_unit': 'm³',
        'geography': 'ESP',
        'year': 2025,
        'month': 1,
        'facility': 'Warehouse_Madrid',
        'description': 'Gas natural para calefacción del almacén'
    },
    {
        'entity_id': 'WAREHOUSE_A',
        'scope': 3,
        'category': 'waste_generated',
        'subcategory': 'landfill',
        'activity_value': 2500,
        'activity_unit': 'kg',
        'geography': 'ESP',
        'year': 2025,
        'month': 1,
        'facility': 'Warehouse_Madrid',
        'description': 'Residuos industriales enviados a vertedero'
    }
]

# Catálogos de categorías por Scope
SCOPE_1_CATEGORIES = [
    "stationary_combustion",
    "mobile_combustion",
    "process_emissions",
    "fugitive_emissions"
]

SCOPE_2_CATEGORIES = [
    "purchased_electricity",
    "purchased_heat",
    "purchased_steam",
    "purchased_cooling"
]

SCOPE_3_CATEGORIES = [
    "purchased_goods_services",
    "capital_goods",
    "fuel_energy_activities",
    "upstream_transportation",
    "waste_generated",
    "business_travel",
    "employee_commuting",
    "upstream_leased_assets",
    "downstream_transportation",
    "processing_sold_products",
    "use_sold_products",
    "end_of_life_sold_products",
    "downstream_leased_assets",
    "franchises",
    "investments"
]

COMMON_UNITS = [
    "litres", "m³", "kg", "tonnes", "kWh", "MWh", "km", 
    "passenger.km", "tonne.km", "USD", "EUR", "GBP"
]

def create_template_excel(output_file='data/PLANTILLA_HUELLA_CARBONO.xlsx'):
    """Crea archivo Excel con múltiples pestañas: Instrucciones, Plantilla, Ejemplos, Catálogos."""
    
    # Crear diccionario de DataFrames para cada hoja
    dfs = {}
    
    # =========================================
    # HOJA 1: INSTRUCCIONES
    # =========================================
    instructions = [
        ['', ''],
        ['PLANTILLA DE CARGA DE DATOS - HUELLA DE CARBONO GHG PROTOCOL', ''],
        ['', ''],
        ['INSTRUCCIONES DE USO:', ''],
        ['', ''],
        ['1. PREPARACIÓN', 'Complete la hoja "DATOS" con las actividades de su organización'],
        ['', 'Utilice los ejemplos de la hoja "EJEMPLOS" como referencia'],
        ['', ''],
        ['2. CAMPOS OBLIGATORIOS (*):', ''],
        ['   - entity_id', 'Identificador único de la entidad emisora (ej: FACTORY001, OFFICE_HQ)'],
        ['   - scope', 'Alcance GHG Protocol: 1 (directo), 2 (indirecto energía), 3 (otros indirectos)'],
        ['   - category', 'Categoría de emisión (ver catálogo en hoja CATEGORÍAS)'],
        ['   - activity_value', 'Valor numérico de la actividad (solo números positivos)'],
        ['   - activity_unit', 'Unidad de medida (ver catálogo en hoja UNIDADES)'],
        ['', ''],
        ['3. CAMPOS OPCIONALES:', ''],
        ['   - subcategory', 'Subcategoría para mayor detalle (ej: diesel, natural_gas, flights)'],
        ['   - geography', 'Código ISO-3 del país (GBR=Reino Unido, USA=Estados Unidos, ESP=España)'],
        ['   - year', 'Año de la actividad (ej: 2025)'],
        ['   - month', 'Mes (1-12)'],
        ['   - facility', 'Nombre de la instalación o sitio'],
        ['   - description', 'Descripción adicional de la actividad'],
        ['', ''],
        ['4. ALCANCES GHG PROTOCOL:', ''],
        ['   Scope 1', 'Emisiones DIRECTAS de fuentes controladas por la organización'],
        ['', '(combustión estacionaria, móvil, procesos, fugas de refrigerantes)'],
        ['   Scope 2', 'Emisiones INDIRECTAS por compra de energía'],
        ['', '(electricidad, calor, vapor, refrigeración adquiridos)'],
        ['   Scope 3', 'Otras emisiones INDIRECTAS en la cadena de valor'],
        ['', '(viajes, residuos, transporte, productos/servicios adquiridos, etc.)'],
        ['', ''],
        ['5. VALIDACIONES AUTOMÁTICAS:', ''],
        ['   ✓', 'El sistema valida que scope sea 1, 2 o 3'],
        ['   ✓', 'Las categorías deben coincidir con el catálogo según el scope'],
        ['   ✓', 'Los valores de actividad deben ser números positivos'],
        ['   ✓', 'Las unidades deben ser compatibles con los factores de emisión'],
        ['', ''],
        ['6. CARGA DEL ARCHIVO:', ''],
        ['   - Formato', 'Guarde el archivo como .xlsx o .csv'],
        ['   - Tamaño', 'Máximo recomendado: 10,000 actividades por archivo'],
        ['   - Encoding', 'UTF-8 para caracteres especiales'],
        ['', ''],
        ['7. SOPORTE:', ''],
        ['   - Documentación', 'Ver README_COMPLETE.md'],
        ['   - Ejemplos', 'Ver hoja EJEMPLOS y carpeta examples/'],
        ['   - Catálogos', 'Ver hojas CATEGORÍAS y UNIDADES'],
        ['', ''],
        ['FECHA DE CREACIÓN:', datetime.now().strftime('%Y-%m-%d %H:%M')],
        ['VERSIÓN:', '1.0.0'],
    ]
    dfs['INSTRUCCIONES'] = pd.DataFrame(instructions, columns=['Campo', 'Descripción'])
    
    # =========================================
    # HOJA 2: DATOS (vacía con headers y descripciones)
    # =========================================
    headers = list(COLUMNS.keys())
    descriptions = list(COLUMNS.values())
    
    # Crear DataFrame vacío con headers
    datos_df = pd.DataFrame(columns=headers)
    # Añadir fila de descripciones
    datos_df.loc[0] = descriptions
    
    dfs['DATOS'] = datos_df
    
    # =========================================
    # HOJA 3: EJEMPLOS
    # =========================================
    ejemplos_df = pd.DataFrame(EXAMPLE_DATA)
    dfs['EJEMPLOS'] = ejemplos_df
    
    # =========================================
    # HOJA 4: CATEGORÍAS
    # =========================================
    max_len = max(len(SCOPE_1_CATEGORIES), len(SCOPE_2_CATEGORIES), len(SCOPE_3_CATEGORIES))
    
    categorias_data = {
        'SCOPE 1 (Emisiones Directas)': SCOPE_1_CATEGORIES + [''] * (max_len - len(SCOPE_1_CATEGORIES)),
        'SCOPE 2 (Energía Indirecta)': SCOPE_2_CATEGORIES + [''] * (max_len - len(SCOPE_2_CATEGORIES)),
        'SCOPE 3 (Cadena de Valor)': SCOPE_3_CATEGORIES + [''] * (max_len - len(SCOPE_3_CATEGORIES))
    }
    dfs['CATEGORÍAS'] = pd.DataFrame(categorias_data)
    
    # =========================================
    # HOJA 5: UNIDADES
    # =========================================
    unidades_data = {
        'Unidad': COMMON_UNITS,
        'Descripción': [
            'Litros (combustibles líquidos)',
            'Metros cúbicos (gas natural)',
            'Kilogramos (residuos, refrigerantes)',
            'Toneladas (materiales, transporte)',
            'Kilovatios-hora (electricidad)',
            'Megavatios-hora (electricidad)',
            'Kilómetros (distancia viajes)',
            'Pasajero-kilómetro (transporte pasajeros)',
            'Tonelada-kilómetro (transporte carga)',
            'Dólares estadounidenses (bienes/servicios)',
            'Euros (bienes/servicios)',
            'Libras esterlinas (bienes/servicios)'
        ]
    }
    dfs['UNIDADES'] = pd.DataFrame(unidades_data)
    
    # =========================================
    # HOJA 6: PAÍSES
    # =========================================
    paises_data = {
        'Código ISO-3': ['GBR', 'USA', 'ESP', 'FRA', 'DEU', 'ITA', 'MEX', 'BRA', 'CHN', 'IND', 'JPN', 'AUS'],
        'País': ['Reino Unido', 'Estados Unidos', 'España', 'Francia', 'Alemania', 'Italia', 
                 'México', 'Brasil', 'China', 'India', 'Japón', 'Australia']
    }
    dfs['PAÍSES'] = pd.DataFrame(paises_data)
    
    # =========================================
    # GUARDAR A EXCEL
    # =========================================
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        for sheet_name, df in dfs.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)
    
    # =========================================
    # APLICAR FORMATO CON OPENPYXL
    # =========================================
    wb = load_workbook(output_file)
    
    # Colores
    header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    example_fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')
    instruction_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    
    header_font = Font(color='FFFFFF', bold=True, size=11)
    title_font = Font(bold=True, size=14, color='1F4E78')
    
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Formatear INSTRUCCIONES
    ws_inst = wb['INSTRUCCIONES']
    ws_inst.column_dimensions['A'].width = 30
    ws_inst.column_dimensions['B'].width = 80
    
    # Título
    ws_inst['A2'].font = title_font
    ws_inst.merge_cells('A2:B2')
    ws_inst['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws_inst.row_dimensions[2].height = 25
    
    # Secciones importantes
    for row in ws_inst.iter_rows(min_row=1, max_row=ws_inst.max_row):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                if cell.value.startswith(('1.', '2.', '3.', '4.', '5.', '6.', '7.')):
                    cell.font = Font(bold=True, size=12, color='1F4E78')
                    cell.fill = instruction_fill
    
    # Formatear DATOS
    ws_datos = wb['DATOS']
    for col_num, column in enumerate(ws_datos.columns, 1):
        ws_datos.column_dimensions[chr(64 + col_num)].width = 20
        
        # Header
        header_cell = ws_datos.cell(row=1, column=col_num)
        header_cell.fill = header_fill
        header_cell.font = header_font
        header_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        header_cell.border = border
        
        # Descripción (fila 2)
        desc_cell = ws_datos.cell(row=2, column=col_num)
        desc_cell.fill = example_fill
        desc_cell.alignment = Alignment(wrap_text=True, vertical='top')
        desc_cell.font = Font(italic=True, size=9)
    
    ws_datos.row_dimensions[1].height = 30
    ws_datos.row_dimensions[2].height = 60
    ws_datos.freeze_panes = 'A3'  # Congelar headers
    
    # Formatear EJEMPLOS
    ws_ejemplos = wb['EJEMPLOS']
    for col_num in range(1, len(ejemplos_df.columns) + 2):
        ws_ejemplos.column_dimensions[chr(64 + col_num)].width = 20
        
        header_cell = ws_ejemplos.cell(row=1, column=col_num)
        header_cell.fill = header_fill
        header_cell.font = header_font
        header_cell.alignment = Alignment(horizontal='center', vertical='center')
        header_cell.border = border
    
    ws_ejemplos.freeze_panes = 'A2'
    
    # Formatear CATEGORÍAS
    ws_cat = wb['CATEGORÍAS']
    for col_num in range(1, 4):
        ws_cat.column_dimensions[chr(64 + col_num)].width = 35
        
        header_cell = ws_cat.cell(row=1, column=col_num)
        header_cell.fill = header_fill
        header_cell.font = header_font
        header_cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        header_cell.border = border
    
    # Formatear UNIDADES
    ws_units = wb['UNIDADES']
    ws_units.column_dimensions['A'].width = 20
    ws_units.column_dimensions['B'].width = 50
    
    for col_num in range(1, 3):
        header_cell = ws_units.cell(row=1, column=col_num)
        header_cell.fill = header_fill
        header_cell.font = header_font
        header_cell.alignment = Alignment(horizontal='center', vertical='center')
        header_cell.border = border
    
    # Formatear PAÍSES
    ws_paises = wb['PAÍSES']
    ws_paises.column_dimensions['A'].width = 20
    ws_paises.column_dimensions['B'].width = 30
    
    for col_num in range(1, 3):
        header_cell = ws_paises.cell(row=1, column=col_num)
        header_cell.fill = header_fill
        header_cell.font = header_font
        header_cell.alignment = Alignment(horizontal='center', vertical='center')
        header_cell.border = border
    
    # Guardar cambios
    wb.save(output_file)
    
    print(f"✅ Plantilla Excel creada exitosamente: {output_file}")
    print(f"\n📊 Hojas incluidas:")
    print(f"   1. INSTRUCCIONES - Guía completa de uso")
    print(f"   2. DATOS - Plantilla vacía para completar")
    print(f"   3. EJEMPLOS - {len(EXAMPLE_DATA)} ejemplos de actividades")
    print(f"   4. CATEGORÍAS - Catálogo completo por Scope")
    print(f"   5. UNIDADES - Unidades de medida soportadas")
    print(f"   6. PAÍSES - Códigos ISO-3 principales")
    print(f"\n🎯 Próximo paso: Complete la hoja 'DATOS' y cargue el archivo en la aplicación")
    
    return output_file


if __name__ == "__main__":
    create_template_excel()
